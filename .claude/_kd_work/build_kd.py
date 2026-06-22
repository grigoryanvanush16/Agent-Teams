# -*- coding: utf-8 -*-
"""Сборка «КД метрики v1 (формат v58)» по эталону ИБ v58 / образцу Маркетинг v1."""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import kd_data as D
import kd_data2 as D2

OUT = r'C:\Users\User\Downloads\КД метрики v1 (формат v58).xlsx'

C_TITLE='1F4E78'; C_SUB='555555'; C_HEAD='1F4E78'; C_BAND='2E75B6'; C_GROUP='D6E4F0'
C_GREEN='C6EFCE'; C_YELLOW='FFEB9C'; C_BLUE='BDD7EE'; C_GRAY='D9D9D9'
C_BI='DDEBF7'; C_REBUILD='FFF2CC'; C_AUTO='E2EFDA'; C_ZEBRA='F4F7FB'; WHITE='FFFFFF'
thin=Side(style='thin', color='D0D0D0'); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(c): return PatternFill('solid', fgColor=c)

READY_FILL={'Считаем сейчас':C_GREEN,'После новой витрины':C_YELLOW,'CH Кирилла — проверить':C_BLUE,'Нет данных (блокер)':C_GRAY}
IMPL_FILL={'BI-сторона':C_BI,'пересобрать':C_REBUILD,'авто в DWH':C_AUTO,'оставить в PBI':C_BI}

wb=openpyxl.Workbook(); wb.remove(wb.active)

def setcol(ws,w):
    for c,v in w.items(): ws.column_dimensions[c].width=v
def title(ws,t,sub,span):
    ws.merge_cells(f'A1:{get_column_letter(span)}1'); ws.merge_cells(f'A2:{get_column_letter(span)}2')
    a=ws.cell(1,1,t); a.font=Font(bold=True,size=14,color=C_TITLE)
    b=ws.cell(2,1,sub); b.font=Font(size=10,color=C_SUB); b.alignment=Alignment(wrap_text=True,vertical='top')
    ws.row_dimensions[2].height=46
def header(ws,row,cols):
    for i,h in enumerate(cols,1):
        c=ws.cell(row,i,h); c.font=Font(bold=True,size=11,color=WHITE); c.fill=fill(C_HEAD)
        c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=BORDER
def band(ws,row,text,span,color=C_BAND,size=12,fc=WHITE):
    ws.merge_cells(f'A{row}:{get_column_letter(span)}{row}')
    c=ws.cell(row,1,text); c.font=Font(bold=True,size=size,color=fc)
    c.fill=fill(color); c.alignment=Alignment(wrap_text=True,vertical='center',indent=1)
def group_cell(ws,row,text):
    c=ws.cell(row,1,text); c.font=Font(bold=True,color=C_TITLE); c.fill=fill(C_GROUP)
    c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=BORDER
def datarow(ws,row,vals,bold_first=True,sql_cols=(),center=()):
    for i,v in enumerate(vals,1):
        c=ws.cell(row,i,v); sz=9 if i in sql_cols else 11
        c.font=Font(size=sz,bold=(bold_first and i==1)); c.border=BORDER
        h='center' if i in center else 'general'
        c.alignment=Alignment(wrap_text=True,vertical='top',horizontal=h)
def sqlblock(ws,row,text,span):
    ws.merge_cells(f'A{row}:{get_column_letter(span)}{row}')
    c=ws.cell(row,1,text); c.font=Font(size=9,name='Consolas'); c.alignment=Alignment(wrap_text=True,vertical='top')
    nlines=text.count(chr(10))+1; ws.row_dimensions[row].height=min(14*nlines+6, 600)

# ===== Лист 1: Каталог метрик =====
# Структура (MECE, 4 блока с подэтапами воронки). Контент берём из D.CATALOG по имени метрики.
NEW_STRUCT=[
 ('I. Воронка КД',[
   ('Привлечение',['Лиды КД','Тип лида (структура лидов)','Факт лиды, ИБ+Аут, загружены в АО']),
   ('Дозвон и качество базы',['Дозвон, шт','% дозвона','Брак лида - недозвон, шт/%','Брак - не давал согласия / негатив, шт/%','Не целевой - ОПФ, шт']),
   ('Квалификация',['Квалифицированные лиды (валовые), шт','% квалификации (валовое)','Квалифицированные лиды (мес в мес), шт','% квалификации (мес в мес)']),
   ('Счета и продажи',['Выставленные счета, шт','Выставленные счета, руб','Продажи тарифы (по дате лида), шт','Продажи итого (по дате лида), шт','Продажи тарифы (по дате оплаты), шт','Продажи итого (по дате оплаты), шт','Продажи разовых услуг, шт','Конверсия по дате лида, %','Конверсия по дате оплаты, %']),
   ('Деньги (выручка, чек, скидки)',['Выручка тарифы (по дате лида), руб','Выручка итого (по дате лида), руб','Выручка тарифы (по дате оплаты), руб','Выручка итого (по дате оплаты), руб','Выручка с разовых услуг, руб','Средний чек тарифы, руб','Средний чек (без разовых Аут и опций ИБ), руб','Ср. кол-во оплаченных месяцев (тарифы), мес','ARPU','Скидка (по дате лида), руб/%','Скидка (по дате оплаты), руб/%']),
 ]),
 ('II. Юнит-экономика КД',[
   ('Расходы и стоимость',['Переменные расходы, руб','ФОТ КД, руб','CPL','CPLq','CPO','CAC']),
 ]),
 ('III. Удержание и LTV',[
   ('Удержание',['Churn (отток), %','% переподписки','Upsell на АУТ, шт/руб','Upsell Сбербанк (WL), руб/шт']),
   ('LTV и маржа',['LTV','Gross margin']),
 ]),
 ('IV. Отчётность (план/факт)',[
   ('План / факт',['Факт на дату','План на дату','% выполнения плана на дату','Пред. год на дату (YoY), руб/%','Пред. мес. на дату (MoM), руб/%']),
 ]),
]
TOP10=['Лиды КД','% дозвона','% квалификации (мес в мес)','Продажи тарифы (по дате оплаты), шт',
       'Конверсия по дате оплаты, %','Выручка тарифы (по дате оплаты), руб','Средний чек тарифы, руб','CPL','CPO','% выполнения плана на дату']
YELLOW_METRICS={'Брак лида - недозвон, шт/%','Переменные расходы, руб','CPL','CPO','Upsell на АУТ, шт/руб','LTV','Gross margin'}
C_YELHL='FFF2CC'
LOOKUP={rec[2]:rec for rec in D.CATALOG if rec[0]!=D.BAND}

STATUS_FILL={'Готово':'C6EFCE','В работе':'FFEB9C','Уточнить':'FCE4D6','GAP':'D9D9D9'}
VIT_FILL={'новая витрина':'FFF2CC','текущая витрина':'E2EFDA','нет данных':'D9D9D9'}
TOP_FILL='FFD966'
# новая витрина = нет готовой таблицы НИ В ОДНОЙ БД (расходы/план - внешний файл; скидки/конверсия - агрегат)
NOVA_VIT={'Переменные расходы, руб','ФОТ КД, руб','CPL','CPLq','CPO','CAC',
          'Скидка (по дате лида), руб/%','Скидка (по дате оплаты), руб/%','План на дату','% выполнения плана на дату',
          'LTV'}                                                    # LTV: LT нашёлся в renewal_detail - строим
NODATA_VIT={'Gross margin'}                                         # true GAP - нет источника себестоимости
# ретеншн переиспользует backoffice.renewal_summary - больше не GAP
STATUS_OVERRIDE={'Churn (отток), %':'Готово','% переподписки':'Готово',
                 'Upsell на АУТ, шт/руб':'Готово','Upsell Сбербанк (WL), руб/шт':'Готово'}
def derive(rec):
    name,status,ready=rec[2],rec[9],rec[10]
    st={'Считаем сейчас':'Готово','После новой витрины':'В работе','Нет данных (блокер)':'GAP'}.get(ready,'В работе')
    if status in ('GAP','Уточнить'): st=status
    if name in STATUS_OVERRIDE: st=STATUS_OVERRIDE[name]
    if name in NODATA_VIT: vit='нет данных'
    elif name in NOVA_VIT: vit='новая витрина'
    else: vit='текущая витрина'
    return st,vit

ws=wb.create_sheet('Каталог метрик')
setcol(ws,{'A':20,'B':4,'C':28,'D':15,'E':45,'F':44,'G':8,'H':6,'I':5,'J':5,'K':11,'L':15})
title(ws,'КД метрики v1 (формат v58) — Каталог метрик',
  'Метрики Коммерческой дирекции по этапам воронки (I), юнит-экономике (II), удержанию и LTV (III) и отчётности (IV). ★ - Топ-10 метрик первой очереди. «Статус» и «Витрина» - цветовая легенда ниже.',12)
# Цветовая легенда (каждое значение - в своей цветной ячейке с расшифровкой)
def legcell(rng,text,col):
    ws.merge_cells(rng); c=ws[rng.split(':')[0]]
    c.value=text; c.fill=fill(col); c.font=Font(size=9,bold=True,color='1F4E78')
    c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); c.border=BORDER
ws['A3']='Статус:'; ws['A3'].font=Font(bold=True,size=9)
legcell('B3:C3','Готово\nсчитаем сейчас',STATUS_FILL['Готово'])
legcell('D3:E3','В работе\nпосле новой витрины',STATUS_FILL['В работе'])
legcell('F3:G3','Уточнить\nоткрытый вопрос',STATUS_FILL['Уточнить'])
legcell('H3:I3','GAP\nнет данных (блокер)',STATUS_FILL['GAP'])
legcell('J3:K3','★ Топ-10\nпервая очередь',TOP_FILL)
ws['A4']='Витрина:'; ws['A4'].font=Font(bold=True,size=9)
legcell('B4:D4','текущая витрина\nготовая таблица - переиспользуем',VIT_FILL['текущая витрина'])
legcell('E4:G4','новая витрина\nстроим (лист «Витрины и таблицы»)',VIT_FILL['новая витрина'])
legcell('H4:K4','нет данных\nблокер источника (Gross margin)',VIT_FILL['нет данных'])
ws.row_dimensions[3].height=26; ws.row_dimensions[4].height=26
header(ws,5,['Группа','№','Метрика','Гранулярность','Описание','Формула (методология словами)','Ед.','Приор.','KPI','★','Статус','Витрина'])
ws.freeze_panes='C6'
r=6; n=0
for block_title,substages in NEW_STRUCT:
    band(ws,r,block_title,12); r+=1
    for substage,names in substages:
        first=True
        for nm in names:
            rec=LOOKUP[nm]; n+=1
            st,vit=derive(rec)
            top='★' if nm in TOP10 else ''
            grp=substage if first else ''
            datarow(ws,r,[grp,n,rec[2],rec[3],rec[4],rec[5],rec[6],rec[7],rec[8],top,st,vit],center=(2,7,8,9,10))
            if grp: group_cell(ws,r,grp)
            if st in STATUS_FILL: ws.cell(r,11).fill=fill(STATUS_FILL[st])
            if vit in VIT_FILL: ws.cell(r,12).fill=fill(VIT_FILL[vit])
            if top:
                ws.cell(r,10).fill=fill(TOP_FILL); ws.cell(r,2).fill=fill(TOP_FILL)
            if nm in YELLOW_METRICS:
                for cc in (3,6): ws.cell(r,cc).fill=fill(C_YELHL)
            first=False; r+=1
band(ws,r,'Жёлтым выделены метрики, уточнённые в «Описание метрик_1»: CPL/CPO = (КВ партнёра + Сертификаты) без прочих и ФОТ; Upsell в АУТ = переход ИБ→АУТ (шт/руб); LTV = ARPU × LT; Gross margin = валовая прибыль; Брак-недозвон - уточнить статусы. Термины: КВ партнёра - комиссионное вознаграждение партнёра; Сертификаты - статья переменных расходов; LT - срок жизни клиента (мес).',12,color=C_YELHL,size=9,fc=C_TITLE)
ws.row_dimensions[r].height=44

# ===== Лист 2: Roadmap витрин =====
ws=wb.create_sheet('Roadmap витрин')
setcol(ws,{'A':26,'B':24,'C':6,'D':14,'E':42,'F':52,'G':30})
title(ws,'Roadmap витрин КД',
  'Что переиспользуем из готовых таблиц (зелёные) и что строим заново (жёлтые). Колонки: витрина/таблица, статус, приоритет, ответственный, источник, зачем нужна, частота. Проверено по marts / owox / backoffice.',7)
n_metrics=sum(1 for x in D.CATALOG if x[0]!=D.BAND)
n_reuse=sum(1 for x in D2.ROADMAP if x[5] in ('переиспользовать','BI-сторона'))
n_build=sum(1 for x in D2.ROADMAP if x[5]=='строить')
for i,(lbl,val) in enumerate([('Итого метрик в каталоге',n_metrics),
        ('Переиспользуем готовых таблиц',n_reuse),('Строим заново витрин',n_build),
        ('Из них ядро (приоритет 0)','dm_kd_funnel')]):
    ws.cell(4+i,1,lbl).font=Font(bold=True); ws.cell(4+i,2,val)
header(ws,9,['Витрина / таблица','Статус','Приор.','Ответственный','Источник','Зачем нужна / что питает','Частота обновления'])
ws.freeze_panes='A10'; r=10
ROAD_FILL={'переиспользовать':'E2EFDA','строить':'FFF2CC','BI-сторона':'DDEBF7'}
for v,st,pr,ow,src,impl,purp in D2.ROADMAP:
    freq={'переиспользовать':'наследует от источника','BI-сторона':'BI-сторона (по источнику)','строить':'Раз в сутки (инкремент)'}.get(impl,'-')
    datarow(ws,r,[v,st,pr,ow,src,purp,freq],center=(3,))
    if impl in ROAD_FILL: ws.cell(r,2).fill=fill(ROAD_FILL[impl])
    r+=1

# ===== Лист 3: CTE =====
ws=wb.create_sheet('CTE'); setcol(ws,{'A':150})
title(ws,'CTE — бизнес-правила и сборка ядра воронки КД',
  'Логика сборки воронки КД: правила связки лид→дозвон→квалификация→продажа→выручка и полный SQL ядра dm_kd_funnel.',1)
sqlblock(ws,4,D2.CTE_RULES,1); r=5
ws.merge_cells('A6:A6'); band(ws,6,D2.CTE_SQL_TITLE,1); sqlblock(ws,7,D2.CTE_SQL,1)

# ===== Лист 4: Витрины и таблицы =====
ws=wb.create_sheet('Витрины и таблицы')
setcol(ws,{'A':24,'B':30,'C':40,'D':15,'E':6,'F':34,'G':38})
title(ws,'Витрины и таблицы (структура по полям)',
  'Подробно описаны НОВЫЕ витрины, которые строим (dm_kd_funnel, dim_kd_cost, plan_kd_monthly, dm_kd_cohort_conversion + тонкие надстройки). Переиспользуемые готовые таблицы (_32_lead_info, _33_opportunities, _23_sale, _44_bill, renewal_summary…) - см. лист «Roadmap витрин». Колонки «вопросы» и «sql» - как в эталоне ИБ v58.',7)
import re as _re
r=4
for vit in D2.VITRINY:
    band(ws,r,vit['title'],7); r+=1
    c=ws.cell(r,1,vit['table']); c.font=Font(bold=True,color=C_TITLE); r+=1
    sqlblock(ws,r,vit['sql'],7); r+=1
    header(ws,r,['поле','описание','комментарий','значения (пример)','ключ','вопросы','sql']); r+=1
    m=_re.search(r'Таблица:\s*([A-Za-z0-9_]+)', vit['table']); tbl=m.group(1) if m else ''
    qmap=D2.QUESTIONS.get(tbl,{})
    for f in vit['fields']:
        name,desc,comm,ex,key,sqlv=f
        q=qmap.get(name,'')
        datarow(ws,r,[name,desc,comm,ex,key,q,sqlv],center=(5,))
        if q: ws.cell(r,6).fill=fill(C_YELHL)
        r+=1
    r+=1

# ===== Лист 5: ТЗ для разработки =====
ws=wb.create_sheet('ТЗ для разработки')
setcol(ws,{'A':4,'B':26,'C':16,'D':24,'E':44,'F':30,'G':36})
title(ws,'ТЗ для разработки (DE)',
  'Задачи для команды данных. «Витрина есть?» - текущее состояние. Каждая задача разблокирует метрики из листа «Каталог метрик». Порядок = приоритет фундамента.',7)
header(ws,4,['#','Задача','Ответственный','Витрина есть?','Что построить','Какие метрики разблокирует','Источник / референс'])
ws.freeze_panes='A5'; r=5
for rec in D2.TZ:
    datarow(ws,r,list(rec),center=(1,)); r+=1

# ===== Лист 6: Реестр дашбордов =====
ws=wb.create_sheet('Реестр дашбордов')
setcol(ws,{'A':30,'B':9,'C':22,'D':6,'E':38,'F':40,'G':34,'H':15})
title(ws,'Реестр дашбордов КД',
  'Дашборды КД и смежные (PBIRS). Источники извлечены из Power Query дашборда «Сквозная отчётность». «Реализация в DWH»: авто в DWH / пересобрать / BI-сторона.',8)
header(ws,4,['Дашборд','Движок','Папка','Исп.','Назначение','Ключевые метрики','Источники данных','Реализация в DWH'])
ws.freeze_panes='A5'; r=5
for d in D2.DASH:
    datarow(ws,r,list(d),center=(4,))
    if d[7] in IMPL_FILL: ws.cell(r,8).fill=fill(IMPL_FILL[d[7]])
    r+=1

# ===== Лист 7: Топ-10 дашбордов =====
ws=wb.create_sheet('Топ-10 дашбордов')
setcol(ws,{'A':6,'B':32,'C':6,'D':9,'E':9,'F':6,'G':44,'H':32,'I':24,'J':16})
title(ws,'Топ-10 дашбордов КД для реформы',
  'Из реестра выбраны опорные дашборды для перевода на DWH. Балл = сумма трёх оценок (1-3): «Исп.» - частота просмотра; «Ценность» - влияние на решения; «Готовн.» - доступность данных в DWH-источниках. Действие: «пересобрать» - свести разрозненные меры; «авто в DWH» - перенести логику почти как есть; «оставить в PBI» - считается на BI поверх источника.',10)
header(ws,4,['Ранг','Дашборд','Исп.','Ценность','Готовн.','Балл','Почему в топе','Метрики','Витрины-источники','Действие'])
ws.freeze_panes='A5'; r=5
for rec in D2.TOP10:
    rank,name,u,v,g,sc,why,met,vit,act=rec
    datarow(ws,r,[rank,name,u,v,g,sc,why,met,vit,act],center=(1,3,4,5,6))
    if act in IMPL_FILL: ws.cell(r,10).fill=fill(IMPL_FILL[act])
    r+=1
band(ws,r,D2.TOP10_NOTE,10,color=C_GROUP,size=10,fc=C_TITLE)

wb.save(OUT)
print('SAVED:', OUT)
print('Метрик в каталоге:', n_metrics, '| переиспользуем:', n_reuse, '| строим:', n_build)
