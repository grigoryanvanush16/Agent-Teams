# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
p=r'C:\Users\User\Downloads\КД метрики v1 (формат v58).xlsx'
wb=openpyxl.load_workbook(p)
print('SHEETS:', wb.sheetnames)
for ws in wb.worksheets:
    print('='*50, ws.title, '| rows', ws.max_row, '| freeze', ws.freeze_panes)
# spot check каталог first data rows
ws=wb['Каталог метрик']
for r in [4,5,6,7,38,42,50,53,54]:
    vals=[ws.cell(r,c).value for c in range(1,14)]
    print('КАТ r%d:'%r, ' | '.join(str(v)[:18] if v is not None else '' for v in vals))
# CTE sheet sanity
ws=wb['CTE']
for r in range(1,9):
    v=ws.cell(r,1).value
    print('CTE r%d:'%r, (str(v)[:70] if v else ''))
# Витрины headers presence
ws=wb['Витрины и таблицы']
bands=[ws.cell(r,1).value for r in range(1,ws.max_row+1) if str(ws.cell(r,1).value).startswith('Витрина')]
print('ВИТРИНЫ:', len(bands))
for b in bands: print('  -', b[:60])
wb.close()
