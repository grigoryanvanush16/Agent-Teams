# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
p = r'C:\Users\User\Downloads\Описание метрик_1 [Keo4DY].xlsx'
wb = openpyxl.load_workbook(p)
ws = wb['Лист1']
def yellow(c):
    f=c.fill
    return bool(f and f.patternType=='solid' and isinstance(f.fgColor.rgb,str) and f.fgColor.rgb[-6:].upper()=='FFFF00')
out=[]
for r in range(1, ws.max_row+1):
    rowcells=[]
    any_yellow=False
    for cc in range(1, ws.max_column+1):
        cell=ws.cell(r,cc)
        v=cell.value
        if v is None: continue
        y=yellow(cell)
        if y: any_yellow=True
        s=str(v).replace('\n',' / ').strip()
        rowcells.append(('[Y]' if y else '')+f'C{cc}: '+s)
    if rowcells:
        marker='  <<< ЖЁЛТАЯ СТРОКА' if any_yellow else ''
        out.append(f'R{r}{marker}')
        for rc in rowcells:
            out.append('   '+rc)
        out.append('')
open('yellow_content.txt','w',encoding='utf-8').write('\n'.join(out))
print('\n'.join(out))
wb.close()
