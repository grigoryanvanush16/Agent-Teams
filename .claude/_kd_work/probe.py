# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
DL = r'C:\Users\User\Downloads'
import os
files = ['ИБ метрики v58.xlsx','Маркетинг метрики v1 (формат v58).xlsx','Описание метрик КД.xlsx','Дашборды для КД(1).xlsx']
for f in files:
    p = os.path.join(DL, f)
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        print('='*70)
        print(f)
        for ws in wb.worksheets:
            print('  SHEET:', repr(ws.title), '| max_row', ws.max_row, 'max_col', ws.max_column)
        wb.close()
    except Exception as e:
        print(f, 'ERR', repr(e))
