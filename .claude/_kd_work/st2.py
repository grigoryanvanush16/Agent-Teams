# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\User\Downloads\Маркетинг метрики v1 (формат v58).xlsx')
ws = wb['Каталог метрик']
seen=set()
for r in range(6, ws.max_row+1):
    for c in (10,11,12):  # J Статус, K Готовность, L Реализация
        cell=ws.cell(r,c)
        if cell.value is None: continue
        f=cell.fill
        fg=f.fgColor.rgb if f and f.patternType else None
        key=(str(cell.value), fg)
        if key in seen: continue
        seen.add(key)
        print(f'col{c}', repr(str(cell.value)[:28]), '| fill', f.patternType, fg)
wb.close()
