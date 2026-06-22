# -*- coding: utf-8 -*-
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
p = r'C:\Users\User\Downloads\Маркетинг метрики v1 (формат v58).xlsx'
wb = openpyxl.load_workbook(p)
for ws in wb.worksheets:
    print('='*60)
    print('SHEET', repr(ws.title))
    print('freeze:', ws.freeze_panes)
    # column widths
    widths = {k: round(v.width,1) if v.width else None for k,v in ws.column_dimensions.items()}
    print('widths:', widths)
    print('merged:', [str(m) for m in ws.merged_cells.ranges][:20])
    # sample header row styling: find first non-empty row and the table header row
    for r in range(1, min(ws.max_row,12)+1):
        for c in range(1, min(ws.max_column,13)+1):
            cell = ws.cell(r,c)
            if cell.value is not None:
                fill = cell.fill
                fg = fill.fgColor.rgb if fill and fill.patternType else None
                font = cell.font
                print(f'  R{r}C{c}', repr(str(cell.value)[:30]),
                      '| fill', fill.patternType, fg,
                      '| bold', font.bold, 'size', font.sz, 'color', font.color.rgb if font.color else None,
                      '| wrap', cell.alignment.wrap_text, 'valign', cell.alignment.vertical)
                break
wb.close()
